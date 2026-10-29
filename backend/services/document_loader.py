"""
文档加载与解析
支持: .pdf / .txt / .md / .docx
输出: 自定义 Document 列表（带 metadata）
完全不依赖 LangChain
"""
from pathlib import Path
from typing import List

from .text_splitter import Document, split_documents
from backend.config import settings


def _load_pdf(file_path: Path) -> List[Document]:
    from pypdf import PdfReader
    reader = PdfReader(str(file_path))
    docs = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        docs.append(Document(page_content=text, metadata={"page": i + 1}))
    return docs


def _load_docx(file_path: Path) -> List[Document]:
    from docx import Document as DocxDoc
    doc = DocxDoc(str(file_path))
    text = "\n".join(p.text for p in doc.paragraphs)
    return [Document(page_content=text, metadata={})]


def _load_text(file_path: Path) -> List[Document]:
    text = file_path.read_text(encoding="utf-8", errors="ignore")
    return [Document(page_content=text, metadata={})]


_LOADER_MAP = {
    ".pdf": _load_pdf,
    ".txt": _load_text,
    ".md": _load_text,  # Markdown 当纯文本
    ".docx": _load_docx,
}


def load_single_document(file_path: Path) -> List[Document]:
    """加载单个文档并附加 source 元数据"""
    suffix = file_path.suffix.lower()
    if suffix not in _LOADER_MAP:
        raise ValueError(f"不支持的文件格式: {suffix}")

    docs = _LOADER_MAP[suffix](file_path)
    for doc in docs:
        doc.metadata.setdefault("source", file_path.name)
    return docs


def load_and_split_documents(file_paths: List[Path]) -> List[Document]:
    """
    加载多个文档，合并后统一分块
    """
    all_docs: List[Document] = []
    for fp in file_paths:
        try:
            docs = load_single_document(fp)
            all_docs.extend(docs)
        except Exception as e:
            print(f"[WARN] 加载 {fp.name} 失败: {e}")

    if not all_docs:
        return []

    splits = split_documents(
        all_docs,
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
    )

    for i, doc in enumerate(splits):
        doc.metadata["chunk_index"] = i

    return splits
